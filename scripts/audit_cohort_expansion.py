"""Clinical eligibility bounds, without expression/state-change analysis.

Lin counts are clinical upper bounds, NOT CD8-qualified participants.
Sade-Feldman: whitelist the clinical sheet and five columns. Published state
summary sheets must not be loaded by this script. No response-based exclusions.
"""
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "reports/evidence/cohort-expansion"


def save_table(name, rows):
    with (ROOT / "tables" / name).open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


source = json.loads((EVIDENCE / "lin-s1-extraction.json").read_text())
assert hashlib.sha256((ROOT / source["source"]).read_bytes()).hexdigest() == source["source_sha256"]
lin_samples = []
for row_number, row in enumerate(source["table"][1:], start=2):
    patient, sample, timepoint, arm, day, site, matched, _response, spatial = row
    lin_samples.append({
        "cohort": "Lin2026_in_house", "patient_id": patient,
        "sample_id": sample, "timepoint": timepoint.lower(),
        "arm_as_reported": arm, "day_as_reported": int(day),
        "anatomical_site": site.replace("\n", " "),
        "matched_as_reported": matched, "drug_name": "not_reported_in_S1",
        "cd8_count": "", "scrna_presence_verified": False,
        "visium_id_as_reported": spatial,
        "source_pdf_page": 11, "source_table_row": row_number,
    })
assert len({r["sample_id"] for r in lin_samples}) == len(lin_samples)
lin_patients = []
for patient in dict.fromkeys(r["patient_id"] for r in lin_samples):
    rows = [r for r in lin_samples if r["patient_id"] == patient]
    pre = [r for r in rows if r["timepoint"] == "pre"]
    on = sorted((r for r in rows if r["timepoint"] == "on"), key=lambda r: r["day_as_reported"])
    assert len(pre) == 1
    assert len({r["arm_as_reported"] for r in rows}) == 1
    lin_patients.append({
        "cohort": "Lin2026_in_house", "patient_id": patient,
        "arm_as_reported": pre[0]["arm_as_reported"],
        "n_pre": len(pre), "n_on": len(on),
        "first_on_day": on[0]["day_as_reported"] if on else "",
        "first_pair_interval_days": on[0]["day_as_reported"] - pre[0]["day_as_reported"] if on else "",
        "first_pair_pre_id": pre[0]["sample_id"],
        "first_pair_on_id": on[0]["sample_id"] if on else "",
        "clinical_pair": bool(pre and on),
        "matched_as_reported": on[0]["matched_as_reported"] if on else "NA",
        "status": "clinical_candidate_only" if on else "no_on_biopsy_in_S1",
        "pending": "specific_drug;scrna_presence;cd8_counts;histology;library_chemistry;source_overlap",
    })

sf_path = ROOT.parent / "data/SadeFeldman2018/raw/Cell2018_TableS1.xlsx"
workbook = openpyxl.load_workbook(sf_path, read_only=True, data_only=True)
sheet = workbook["Patient-scRNA data"]


def biopsy_day(value):
    value = str(value or "").strip()
    if value in {"", "NA"}:
        return None
    match = re.match(r"\(?\s*(-?\d+)\)?\s*;", value)
    if not match:
        raise ValueError(f"Unexpected biopsy day format: {value!r}")
    return int(match.group(1))


sf_patients = []
for row_number in range(3, sheet.max_row + 1):
    patient = sheet.cell(row_number, 2).value
    if not isinstance(patient, str) or not re.fullmatch(r"P\d+", patient):
        continue
    therapy = str(sheet.cell(row_number, 5).value).strip()
    # No response fields, clinical outcomes, molecular annotations or state sheets.
    pre, on1, on2 = [biopsy_day(sheet.cell(row_number, col).value) for col in [7, 8, 9]]
    sf_patients.append({
        "cohort": "SadeFeldman2018", "patient_id": patient,
        "therapy_as_reported": therapy, "pure_pd1_label": therapy == "PD1",
        "pre_day": pre, "first_on_day": on1, "second_on_day": on2,
        "first_pair_interval_days": on1 - pre if pre is not None and on1 is not None else None,
        "clinical_pair": pre is not None and on1 is not None,
        "specific_drug_verified": False, "cd8_count_verified": False,
        "source_sheet": sheet.title, "source_row": row_number,
        "source_columns": "B;E;G;H;I",
    })
workbook.close()

with (ROOT / "tables/samples.tsv").open() as handle:
    special_samples = list(csv.DictReader(handle, delimiter="\t"))
special_patients = []
for patient in sorted({r["patient_id"] for r in special_samples}):
    rows = [r for r in special_samples if r["patient_id"] == patient]
    present = [r for r in rows if r["present_in_immune_object"].lower() == "true"]
    pre = [int(r["author_annotated_cd8_cells"]) for r in present if r["timepoint"] == "pre"]
    on = [int(r["author_annotated_cd8_cells"]) for r in present if r["timepoint"] == "on"]
    result = {"cohort": "SPECIAL", "patient_id": patient, "arm": rows[0]["arm"],
              "max_pre_cd8": max(pre) if pre else None, "max_on_cd8": max(on) if on else None,
              "interval_verified": False}
    for threshold in [15, 30, 50]:
        result[f"count_only_candidate_{threshold}"] = bool(pre and on and max(pre) >= threshold and max(on) >= threshold)
    special_patients.append(result)

save_table("cohort_expansion_lin_samples.tsv", lin_samples)
save_table("cohort_expansion_lin_patients.tsv", lin_patients)
save_table("cohort_expansion_sade_patients.tsv", sf_patients)
save_table("cohort_expansion_special_by_arm.tsv", special_patients)

lin_mono = [r for r in lin_patients if r["clinical_pair"] and r["arm_as_reported"] == "Mono"]
lin_combo = [r for r in lin_patients if r["clinical_pair"] and r["arm_as_reported"] == "Combo"]
special_mono = [r for r in special_patients if r["count_only_candidate_30"] and r["arm"] == "anti_pd1_monotherapy"]
special_combo = [r for r in special_patients if r["count_only_candidate_30"] and r["arm"] == "combination"]
sf_mono = [r for r in sf_patients if r["clinical_pair"] and r["pure_pd1_label"]]
summary = {
    "status": "necessary_bound_only_no_G1_pass",
    "input_hashes": {str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p): hashlib.sha256(p.read_bytes()).hexdigest()
                     for p in [EVIDENCE / "lin-s1-extraction.json", sf_path, ROOT / "tables/samples.tsv"]},
    "lin": {"clinical_samples": len(lin_samples), "clinical_patients": len(lin_patients),
            "clinical_paired_patients": sum(r["clinical_pair"] for r in lin_patients),
            "mono_pair_ids": [r["patient_id"] for r in lin_mono],
            "combo_pair_ids": [r["patient_id"] for r in lin_combo],
            "mono_first_on_days": sorted(set(r["first_on_day"] for r in lin_mono)),
            "drug_names_verified": False, "cd8_counts_verified": False},
    "sade_feldman": {"clinical_patients": len(sf_patients),
                     "clinical_paired_patients": sum(r["clinical_pair"] for r in sf_patients),
                     "pure_pd1_clinical_pair_ids": [r["patient_id"] for r in sf_mono],
                     "pure_pd1_first_on_days": [r["first_on_day"] for r in sf_mono],
                     "pure_pd1_interbiopsy_days": [r["first_pair_interval_days"] for r in sf_mono],
                     "drug_names_verified": False, "cd8_counts_verified": False,
                     "intact_final_test_claim_allowed": False},
    "special": {"mono_count_candidate_30_ids": [r["patient_id"] for r in special_mono],
                "combo_count_candidate_30_ids": [r["patient_id"] for r in special_combo]},
    "optimistic_bounds_not_eligibility": {
        "early_mono_special_plus_lin": len(special_mono) + len(lin_mono),
        "late_and_early_mono_with_sade": len(special_mono) + len(lin_mono) + len(sf_mono),
        "early_mono_and_combo_special_plus_lin": len(special_mono) + len(lin_mono) + len(special_combo) + len(lin_combo),
        "required_original_development_plus_lockbox": 14,
        "assumptions": "All clinical candidates have paired scRNA and >=30 CD8, allowed drugs, valid intervals and no cross-study duplicates; NONE of these aggregate bounds confirms eligibility.",
    },
    "state_changes_computed": False, "lockbox_assigned": False,
    "published_sade_state_summaries_exposed_during_initial_sheet_inventory": True,
}
(EVIDENCE / "audit.json").write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n")
print(json.dumps(summary, indent=2))
